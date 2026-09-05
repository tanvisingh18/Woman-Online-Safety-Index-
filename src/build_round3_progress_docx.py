"""
Build black-and-white protocol-style DOCX from the Round-3 canonical markdown.

Run:
  PYTHONPATH=src python src/build_round3_progress_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

SRC = Path("docs/FINAL_PROGRESS_REPORT_AUGUST_2026_ROUND3.md")
OUT = Path("docs/Women_Online_Safety_Index_Progress_Report_August_2026_Round3.docx")


def _set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic


def _set_paragraph_spacing(p, before=0, after=6, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def _add_horizontal_line(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_formatted_text(paragraph, text: str, base_size=11):
    """Support **bold** and *italic* and `code` spans lightly."""
    # Split keeping markers
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, size=base_size, bold=True)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, size=base_size, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, size=base_size)
        else:
            # strip residual markdown links [text](url) → text
            clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part)
            clean = clean.replace("\\[", "[").replace("\\]", "]")
            run = paragraph.add_run(clean)
            _set_run_font(run, size=base_size)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h.strip())
        _set_run_font(run, size=10, bold=True)
        # light header shading
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EEEEEE")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tePr = hdr[i]._tc.get_or_add_tcPr()
        hdr[i]._tc.get_or_add_tcPr().append(shading)

    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            # allow bold markers inside cells
            _add_formatted_text(p, val.strip(), base_size=9)

    doc.add_paragraph()  # spacer


def _parse_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    i = start + 1
    # skip separator
    if i < len(lines) and re.match(r"^\|?\s*-+", lines[i]):
        i += 1
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        # pad/truncate
        if len(row) < len(header):
            row += [""] * (len(header) - len(row))
        rows.append(row[: len(header)])
        i += 1
    return header, rows, i


def build() -> Path:
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)

    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        # skip empty
        if not raw.strip():
            i += 1
            continue

        # horizontal rule
        if raw.strip() == "---":
            _add_horizontal_line(doc)
            i += 1
            continue

        # fenced code block
        if raw.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # closing fence
            p = doc.add_paragraph()
            _set_paragraph_spacing(p, before=4, after=8, line=1.0)
            run = p.add_run("\n".join(code_lines))
            _set_run_font(run, size=9)
            continue

        # table
        if raw.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*-+", lines[i + 1]):
            headers, rows, i = _parse_table(lines, i)
            _add_table(doc, headers, rows)
            continue

        # headings
        if raw.startswith("# "):
            p = doc.add_heading("", level=1)
            run = p.add_run(raw[2:].strip())
            _set_run_font(run, size=14, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_paragraph_spacing(p, before=14, after=8)
            i += 1
            continue
        if raw.startswith("## "):
            p = doc.add_heading("", level=2)
            run = p.add_run(raw[3:].strip())
            _set_run_font(run, size=12, bold=True)
            _set_paragraph_spacing(p, before=12, after=6)
            i += 1
            continue
        if raw.startswith("### "):
            p = doc.add_heading("", level=3)
            run = p.add_run(raw[4:].strip())
            _set_run_font(run, size=11, bold=True)
            _set_paragraph_spacing(p, before=10, after=4)
            i += 1
            continue

        # bullet / numbered
        m_bullet = re.match(r"^[-*] (.+)$", raw)
        m_num = re.match(r"^(\d+)\. (.+)$", raw)
        if m_bullet or m_num:
            text_body = m_bullet.group(1) if m_bullet else m_num.group(2)
            style_name = "List Bullet" if m_bullet else "List Number"
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph()
                prefix = "• " if m_bullet else f"{m_num.group(1)}. "
                text_body = prefix + text_body
            _add_formatted_text(p, text_body)
            _set_paragraph_spacing(p, before=0, after=3, line=1.1)
            i += 1
            continue

        # display equation-ish lines (skip raw LaTeX blocks lightly)
        if raw.strip().startswith("\\[") or raw.strip() == "\\]" or raw.strip().startswith("w_i"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(raw.strip().replace("\\[", "").replace("\\]", "").replace("\\text{", "").replace("}", ""))
            _set_run_font(run, size=11, italic=True)
            i += 1
            continue

        # normal paragraph (may continue until blank? keep simple: one md line = one para)
        p = doc.add_paragraph()
        _add_formatted_text(p, raw.strip())
        _set_paragraph_spacing(p, before=0, after=6, line=1.15)
        i += 1

    # Footer note
    _add_horizontal_line(doc)
    p = doc.add_paragraph()
    run = p.add_run(
        "Generated from docs/FINAL_PROGRESS_REPORT_AUGUST_2026_ROUND3.md — "
        "Faculty Round 3 complete (IPW, hotspot quarantine, Telegram lit-adj retired, "
        "Paths A/B/C, Wilson intervals, FP taxonomy). Black-and-white protocol style."
    )
    _set_run_font(run, size=9, italic=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote → {path.resolve()}")
